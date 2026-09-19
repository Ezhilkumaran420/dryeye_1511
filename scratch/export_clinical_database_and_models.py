import openpyxl
import os, glob, re, json
import numpy as np

eye_path = r"C:\Users\ezhil\Downloads\Eye classification.xlsx"
npd_path = r"C:\Users\ezhil\Downloads\N & PD Eye.xlsx"
img_dir = r"D:\AI+ML"

wb_eye = openpyxl.load_workbook(eye_path, data_only=True)
wb_npd = openpyxl.load_workbook(npd_path, data_only=True)

# Parse NPD
npd_rows = list(wb_npd['OSDI_Final'].iter_rows(values_only=True))[1:]
npd_data = {}
for r in npd_rows:
    sno = int(r[0])
    npd_data[sno] = {
        're': str(r[1]).strip(),
        'osdi': float(r[2] or 0),
        'le': str(r[3]).strip()
    }

# Parse Eye Classification
re_rows = list(wb_eye['RE'].iter_rows(values_only=True))[1:]
le_rows = list(wb_eye['LE'].iter_rows(values_only=True))[1:]

clinical_registry = {}

for r in re_rows:
    sno = int(r[0])
    lbl = str(r[-1]).strip()
    is_dry = 1 if 'dry' in lbl.lower() else 0
    clinical_registry[sno] = {
        'sno': sno,
        'osdi': float(r[1] or (npd_data.get(sno, {}).get('osdi', 0))),
        'age': r[2],
        'gender': r[3],
        're': {
            'diagnosis': 'Possible Dry Eye' if is_dry else 'Normal',
            'is_dry': is_dry,
            'cr10': round(float(r[4] or 0), 3),
            'cr7': round(float(r[5] or 0), 3),
            'nc': round(float(r[6] or 0), 2),
            'tc': round(float(r[7] or 0), 2),
            'cc': round(float(r[8] or 0), 2),
            'nl': round(float(r[9] or 0), 2),
            'tl': round(float(r[10] or 0), 2),
            't0': round(float(r[11] or 0), 2),
            't10': round(float(r[12] or 0), 2),
            'most': round(float(r[13] or 0), 2),
        }
    }

for r in le_rows:
    sno = int(r[0])
    lbl = str(r[-1]).strip()
    is_dry = 1 if 'dry' in lbl.lower() else 0
    if sno in clinical_registry:
        clinical_registry[sno]['le'] = {
            'diagnosis': 'Possible Dry Eye' if is_dry else 'Normal',
            'is_dry': is_dry,
            'cr10': round(float(r[2] or 0), 3),
            'cr7': round(float(r[3] or 0), 3),
            'nc': round(float(r[4] or 0), 2),
            'tc': round(float(r[5] or 0), 2),
            'cc': round(float(r[6] or 0), 2),
            'nl': round(float(r[7] or 0), 2),
            'tl': round(float(r[8] or 0), 2),
            't0': round(float(r[9] or 0), 2),
            't10': round(float(r[10] or 0), 2),
            'most': round(float(r[11] or 0), 2),
        }

# Determine overall classification for each subject
for sno, s in clinical_registry.items():
    re_dry = s['re']['is_dry'] == 1
    le_dry = s.get('le', {}).get('is_dry', 0) == 1
    if re_dry and le_dry:
        s['overall_category'] = 'Bilateral Dry Eye'
        s['overall_diagnosis'] = 'Possible Dry Eye'
        s['is_overall_dry'] = 1
    elif not re_dry and not le_dry:
        s['overall_category'] = 'Healthy Normal'
        s['overall_diagnosis'] = 'Normal'
        s['is_overall_dry'] = 0
    elif not re_dry and le_dry:
        s['overall_category'] = 'Asymmetric (RE Normal / LE Dry)'
        s['overall_diagnosis'] = 'Possible Dry Eye'
        s['is_overall_dry'] = 1
    else:
        s['overall_category'] = 'Asymmetric (RE Dry / LE Normal)'
        s['overall_diagnosis'] = 'Possible Dry Eye'
        s['is_overall_dry'] = 1
        
    s['delta_t'] = round(abs(s['re']['cc'] - s['le']['cc']), 2)

# Build training eye samples (156 samples)
samples = []
for sno, s in clinical_registry.items():
    # RE sample
    samples.append({
        'sno': sno,
        'eye': 'RE',
        'y': s['re']['is_dry'],
        'features_full': [
            s['osdi'],
            s['re']['cr10'],
            s['re']['cr7'],
            s['re']['nc'],
            s['re']['tc'],
            s['re']['cc'],
            s['re']['nl'],
            s['re']['tl'],
            s['re']['t0'],
            s['re']['t10'],
            s['re']['most']
        ],
        'features_thermal_only': [
            s['re']['cr10'],
            s['re']['cr7'],
            s['re']['nc'],
            s['re']['tc'],
            s['re']['cc'],
            s['re']['nl'],
            s['re']['tl'],
            s['re']['t0'],
            s['re']['t10'],
            s['re']['most']
        ]
    })
    # LE sample
    samples.append({
        'sno': sno,
        'eye': 'LE',
        'y': s['le']['is_dry'],
        'features_full': [
            s['osdi'],
            s['le']['cr10'],
            s['le']['cr7'],
            s['le']['nc'],
            s['le']['tc'],
            s['le']['cc'],
            s['le']['nl'],
            s['le']['tl'],
            s['le']['t0'],
            s['le']['t10'],
            s['le']['most']
        ],
        'features_thermal_only': [
            s['le']['cr10'],
            s['le']['cr7'],
            s['le']['nc'],
            s['le']['tc'],
            s['le']['cc'],
            s['le']['nl'],
            s['le']['tl'],
            s['le']['t0'],
            s['le']['t10'],
            s['le']['most']
        ]
    })

def train_optimal_model(feat_key):
    X = np.array([s[feat_key] for s in samples])
    y = np.array([s['y'] for s in samples])
    
    mean = np.mean(X, axis=0)
    std = np.std(X, axis=0)
    std[std == 0] = 1.0
    X_norm = (X - mean) / std
    
    N, D = X_norm.shape
    w = np.zeros(D)
    b = 0.0
    
    pos_weight = (len(y) - sum(y)) / sum(y)
    lr = 0.08
    epochs = 2500
    reg = 0.01
    
    for _ in range(epochs):
        z = np.dot(X_norm, w) + b
        p = 1.0 / (1.0 + np.exp(-np.clip(z, -15.0, 15.0)))
        err = (p - y)
        weights = np.where(y == 1, pos_weight, 1.0)
        weighted_err = err * weights
        grad_w = np.dot(X_norm.T, weighted_err) / N + reg * w
        grad_b = np.mean(weighted_err)
        w -= lr * grad_w
        b -= lr * grad_b
        
    z_final = np.dot(X_norm, w) + b
    probs = 1.0 / (1.0 + np.exp(-np.clip(z_final, -15.0, 15.0)))
    
    best_acc, best_thresh = 0, 0.5
    for thresh in np.linspace(0.2, 0.8, 61):
        preds = (probs >= thresh).astype(int)
        acc = np.mean(preds == y)
        tp = np.sum((preds == 1) & (y == 1))
        fn = np.sum((preds == 0) & (y == 1))
        sens = tp / (tp + fn) if (tp + fn) > 0 else 0
        tn = np.sum((preds == 0) & (y == 0))
        fp = np.sum((preds == 1) & (y == 0))
        spec = tn / (tn + fp) if (tn + fp) > 0 else 0
        if acc > best_acc and sens >= 0.80 and spec >= 0.80:
            best_acc = acc
            best_thresh = thresh
            
    preds = (probs >= best_thresh).astype(int)
    tp = int(np.sum((preds == 1) & (y == 1)))
    tn = int(np.sum((preds == 0) & (y == 0)))
    fp = int(np.sum((preds == 1) & (y == 0)))
    fn = int(np.sum((preds == 0) & (y == 1)))
    sens = tp / (tp + fn)
    spec = tn / (tn + fp)
    acc = np.mean(preds == y)
    
    return {
        'weights': [round(float(val), 4) for val in w],
        'bias': round(float(b), 4),
        'mean': [round(float(val), 4) for val in mean],
        'std': [round(float(val), 4) for val in std],
        'threshold': round(float(best_thresh), 3),
        'accuracy': round(float(acc), 4),
        'sensitivity': round(float(sens), 4),
        'specificity': round(float(spec), 4),
        'tp': tp, 'tn': tn, 'fp': fp, 'fn': fn,
        'total_samples': N
    }

model_multimodal = train_optimal_model('features_full')
model_thermal_only = train_optimal_model('features_thermal_only')

print("Multimodal Model:", model_multimodal)
print("\nThermal Only Model:", model_thermal_only)

# Export payload
export_payload = {
    'model_multimodal': model_multimodal,
    'model_thermal_only': model_thermal_only,
    'clinical_registry': clinical_registry
}

with open("scratch/clinical_models_and_registry.json", "w", encoding="utf-8") as f:
    json.dump(export_payload, f, indent=2)

print(f"\nSuccessfully exported clinical registry ({len(clinical_registry)} subjects) and balanced models!")
