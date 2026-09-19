import openpyxl
import os, glob, re, json
import numpy as np

eye_path = r"C:\Users\ezhil\Downloads\Eye classification.xlsx"
npd_path = r"C:\Users\ezhil\Downloads\N & PD Eye.xlsx"
img_dir = r"D:\AI+ML"

wb_eye = openpyxl.load_workbook(eye_path, data_only=True)
wb_npd = openpyxl.load_workbook(npd_path, data_only=True)

# Parse NPD
npd_rows = list(wb_npd['OSDI_Final'].iter_rows(values_only=True))[1:]
npd_data = {}
for r in npd_rows:
    sno = int(r[0])
    npd_data[sno] = {
        're': str(r[1]).strip(),
        'osdi': float(r[2] or 0),
        'le': str(r[3]).strip()
    }

# Parse Eye Classification
re_rows = list(wb_eye['RE'].iter_rows(values_only=True))[1:]
le_rows = list(wb_eye['LE'].iter_rows(values_only=True))[1:]

all_subjects = {}

for r in re_rows:
    sno = int(r[0])
    lbl = str(r[-1]).strip()
    is_dry = 1 if 'dry' in lbl.lower() else 0
    all_subjects[sno] = {
        'sno': sno,
        'osdi': float(r[1] or 0),
        'age': r[2],
        'gender': r[3],
        're': {
            'diagnosis': 'Possible Dry Eye' if is_dry else 'Normal',
            'is_dry': is_dry,
            'cr10': float(r[4] or 0),
            'cr7': float(r[5] or 0),
            'nc': float(r[6] or 0),
            'tc': float(r[7] or 0),
            'cc': float(r[8] or 0),
            'nl': float(r[9] or 0),
            'tl': float(r[10] or 0),
            't0': float(r[11] or 0),
            't10': float(r[12] or 0),
            'most': float(r[13] or 0),
            'n': float(r[14] or 0) if len(r)>14 and r[14] else float(r[6] or 0)-0.38,
            'c': float(r[15] or 0) if len(r)>15 and r[15] else float(r[8] or 0),
            't': float(r[16] or 0) if len(r)>16 and r[16] else float(r[7] or 0)-0.72
        }
    }

for r in le_rows:
    sno = int(r[0])
    lbl = str(r[-1]).strip()
    is_dry = 1 if 'dry' in lbl.lower() else 0
    if sno in all_subjects:
        all_subjects[sno]['le'] = {
            'diagnosis': 'Possible Dry Eye' if is_dry else 'Normal',
            'is_dry': is_dry,
            'cr10': float(r[2] or 0),
            'cr7': float(r[3] or 0),
            'nc': float(r[4] or 0),
            'tc': float(r[5] or 0),
            'cc': float(r[6] or 0),
            'nl': float(r[7] or 0),
            'tl': float(r[8] or 0),
            't0': float(r[9] or 0),
            't10': float(r[10] or 0),
            'most': float(r[11] or 0),
            'n': float(r[12] or 0) if len(r)>12 and r[12] else float(r[4] or 0)-0.38,
            'c': float(r[13] or 0) if len(r)>13 and r[13] else float(r[6] or 0),
            't': float(r[14] or 0) if len(r)>14 and r[14] else float(r[5] or 0)-0.72
        }

# Determine overall classification for each subject
for sno, s in all_subjects.items():
    re_dry = s['re']['is_dry'] == 1
    le_dry = s.get('le', {}).get('is_dry', 0) == 1
    if re_dry and le_dry:
        s['overall_category'] = 'Bilateral Dry Eye'
        s['overall_diagnosis'] = 'Possible Dry Eye'
    elif not re_dry and not le_dry:
        s['overall_category'] = 'Healthy Normal'
        s['overall_diagnosis'] = 'Normal'
    elif not re_dry and le_dry:
        s['overall_category'] = 'Asymmetric (RE Normal / LE Dry)'
        s['overall_diagnosis'] = 'Possible Dry Eye'
    else:
        s['overall_category'] = 'Asymmetric (RE Dry / LE Normal)'
        s['overall_diagnosis'] = 'Possible Dry Eye'

print(f"Total structured subjects: {len(all_subjects)}")
re_dry_cnt = sum(1 for s in all_subjects.values() if s['re']['is_dry'])
le_dry_cnt = sum(1 for s in all_subjects.values() if s.get('le',{}).get('is_dry'))
print(f"Dry Eye counts: RE={re_dry_cnt}/78, LE={le_dry_cnt}/78")
