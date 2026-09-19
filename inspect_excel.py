import openpyxl
import os

def inspect_file(filepath, name):
    print("=" * 70)
    print(f"EXCEL FILE: {name} ({filepath})")
    print("=" * 70)
    if not os.path.exists(filepath):
        print(f"ERROR: File not found: {filepath}")
        return

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        print("Sheet names:", wb.sheetnames)
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            print(f"\n--- Sheet: {sheetname} ---")
            print(f"Dimensions: {ws.max_row} rows x {ws.max_column} columns")
            
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                print("Sheet is empty.")
                continue
            
            headers = rows[0]
            print(f"Headers ({len(headers)}):", headers)
            
            data_rows = rows[1:]
            print(f"Total data rows: {len(data_rows)}")
            print("\nFirst 10 rows:")
            for i, row in enumerate(data_rows[:10]):
                print(f"  Row {i+1}: {row}")
                
            print("\nUnique values per column summary:")
            for col_idx, col_name in enumerate(headers):
                vals = [r[col_idx] for r in data_rows if col_idx < len(r) and r[col_idx] is not None]
                unique_vals = list(set(vals))
                sample = unique_vals[:15]
                print(f"  Col {col_idx+1} [{col_name}]: {len(unique_vals)} unique values, sample: {sample}")
    except Exception as e:
        print("ERROR:", e)
        import traceback; traceback.print_exc()
    print()

inspect_file(r"C:\Users\ezhil\Downloads\Eye classification.xlsx", "Eye classification.xlsx")
inspect_file(r"C:\Users\ezhil\Downloads\N & PD Eye.xlsx", "N & PD Eye.xlsx")
